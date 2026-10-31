"""Чтение исходных таблиц из файлов."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from .header import (
    HEADER_SCAN_ROWS,
    build_headers,
    find_column_number_row,
    find_header_top,
    trim_signature_rows,
)
from .models import TableData
from .normalize import is_empty, normalize_header


def _to_python(value: Any) -> Any:
    if pd.isna(value):
        return None
    if hasattr(value, "to_pydatetime"):
        return value.to_pydatetime()
    return value


def _read_excel_simple(path: Path, sheet: int | str) -> TableData:
    """Читает таблицу с шапкой в первой строке (запасной вариант)."""
    frame = pd.read_excel(path, sheet_name=sheet, header=0, dtype=object)

    headers = [str(column) for column in frame.columns]
    rows: list[list[Any]] = []
    dropped_empty_rows = 0

    for record in frame.itertuples(index=False, name=None):
        row = [_to_python(value) for value in record]
        if all(is_empty(value) for value in row):
            dropped_empty_rows += 1
            continue
        rows.append(row)

    return TableData(
        source=str(path),
        headers=headers,
        rows=rows,
        metadata={
            "sheet": sheet,
            "dropped_empty_rows": dropped_empty_rows,
            "header_row": 1,
        },
    )


def _sample_rows(worksheet: Any, width: int, limit: int) -> list[list[Any]]:
    last_row = min(worksheet.max_row, limit)
    return [
        [worksheet.cell(row=row, column=column).value for column in range(1, width + 1)]
        for row in range(1, last_row + 1)
    ]


def _expanded_header_matrix(
    worksheet: Any,
    top: int,
    end: int,
    width: int,
) -> list[list[Any]]:
    """Копия строк шапки с раскрытыми объединёнными ячейками.

    ``top`` и ``end`` — индексы с нулём, ``end`` не включается.
    """
    matrix = [
        [worksheet.cell(row=row + 1, column=column + 1).value for column in range(width)]
        for row in range(top, end)
    ]

    for merged in worksheet.merged_cells.ranges:
        if merged.max_row <= top or merged.min_row > end:
            continue

        row_from = max(merged.min_row, top + 1) - 1
        row_to = min(merged.max_row, end)
        column_from = merged.min_col - 1
        column_to = min(merged.max_col, width)
        value = worksheet.cell(
            row=merged.min_row,
            column=merged.min_col,
        ).value

        for row in range(row_from, row_to):
            for column in range(column_from, column_to):
                matrix[row - top][column] = value

    return matrix


def _merged_column_continuations(
    worksheet: Any,
    first_row: int,
    last_row: int,
    width: int,
) -> set[int]:
    """Колонки, которые являются продолжением объединённой ячейки."""
    continuations: set[int] = set()
    for merged in worksheet.merged_cells.ranges:
        if merged.max_row < first_row or merged.min_row > last_row:
            continue
        for column in range(merged.min_col + 1, min(merged.max_col, width) + 1):
            continuations.add(column - 1)
    return continuations


def _select_columns(
    headers: list[str],
    continuations: set[int],
) -> tuple[list[int], list[int], list[int]]:
    """Убирает пустые и «схлопнувшиеся» объединённые колонки."""
    keep: list[int] = []
    collapsed: list[int] = []
    empty: list[int] = []

    for index, header in enumerate(headers):
        key = normalize_header(header)
        if not key:
            empty.append(index)
            continue
        if (
            keep
            and index in continuations
            and key == normalize_header(headers[keep[-1]])
        ):
            collapsed.append(index)
            continue
        keep.append(index)

    return keep, collapsed, empty


def read_excel(path: str | Path, sheet: int | str = 0) -> TableData:
    """Читает один лист Excel в TableData.

    Если в начале листа есть преамбула, шапка ищется по характерным
    заголовкам. Поддерживаются многострочная шапка и объединённые ячейки.
    Пустые строки и строки подписей пропускаются.
    """
    excel_path = Path(path).expanduser()
    workbook = load_workbook(excel_path, read_only=False, data_only=True)

    try:
        if isinstance(sheet, int):
            worksheet = workbook.worksheets[sheet]
        else:
            worksheet = workbook[sheet]

        width = worksheet.max_column
        sample = _sample_rows(worksheet, width, HEADER_SCAN_ROWS)
        top = find_header_top(sample)

        if top is None:
            return _read_excel_simple(excel_path, sheet)

        column_numbers = find_column_number_row(sample, top)
        if column_numbers is None:
            header_end = top + 1
            data_start = top + 1
        else:
            header_end = column_numbers
            data_start = column_numbers + 1

        header_matrix = _expanded_header_matrix(
            worksheet,
            top,
            header_end,
            width,
        )
        headers = build_headers(header_matrix, width)
        continuations = _merged_column_continuations(
            worksheet,
            top + 1,
            header_end,
            width,
        )
        keep, collapsed, empty = _select_columns(headers, continuations)

        if not keep:
            return _read_excel_simple(excel_path, sheet)

        warnings: list[str] = []
        if collapsed:
            warnings.append(
                f"Объединённые колонки шапки схлопнуты: {len(collapsed)}."
            )
        if empty:
            warnings.append(f"Колонки без заголовка пропущены: {len(empty)}.")

        headers = [headers[index] for index in keep]
        rows: list[list[Any]] = []
        dropped_empty_rows = 0

        for row_index in range(data_start, worksheet.max_row):
            row = [
                _to_python(
                    worksheet.cell(row=row_index + 1, column=column + 1).value
                )
                for column in keep
            ]
            if all(is_empty(value) for value in row):
                dropped_empty_rows += 1
                continue
            rows.append(row)

        rows, trimmed_rows = trim_signature_rows(rows, warnings)

        return TableData(
            source=str(excel_path),
            headers=headers,
            rows=rows,
            metadata={
                "sheet": sheet,
                "header_row": top + 1,
                "data_start_row": data_start + 1,
                "dropped_empty_rows": dropped_empty_rows,
                "collapsed_columns": collapsed,
                "trimmed_signature_rows": trimmed_rows,
                "warnings": warnings,
            },
        )
    finally:
        workbook.close()
