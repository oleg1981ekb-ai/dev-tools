import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
import datetime
import os

def load_tasks(json_path="tasks.json"):
    if not os.path.exists(json_path):
        return []
    with open(json_path, "r", encoding="utf-8") as f:
        return json.load(f)

def save_tasks(tasks, json_path="tasks.json"):
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(tasks, f, ensure_ascii=False, indent=2)

def get_date_input_with_exit(prompt, default_date):
    while True:
        user_input = input(f"{prompt} (ГГГГ-ММ-СС) [Enter для {default_date}]: ").strip()
        if user_input.lower() in ['выход', 'exit', 'quit']:
            return "exit"
        if not user_input:
            return default_date
        try:
            return datetime.datetime.strptime(user_input, "%Y-%m-%d").date()
        except ValueError:
            print("❌ Неверный формат! Используйте ГГГГ-ММ-СС (например, 2026-10-01).")

def input_tasks_interactively():
    print("\n" + "="*50)
    print("      РЕЖИМ ПОСЛЕДОВАТЕЛЬНОГО ВВОДА РАБОТ")
    print("="*50)
    print("Введите 'выход' в любой момент, чтобы завершить ввод.\n")
    
    new_tasks = []
    last_end_date = None
    
    while True:
        num = input("1. Введите номер пункта (например, 1.1 или 2): ").strip()
        if num.lower() in ['выход', 'exit', 'quit']:
            break
            
        name = input("2. Введите наименование (или название раздела): ").strip()
        if name.lower() in ['выход', 'exit', 'quit']:
            break
            
        is_section_input = input("Это название РАЗДЕЛА? (д/н) [по умолчанию н]: ").strip()
        if is_section_input.lower() in ['выход', 'exit', 'quit']:
            break
        is_section = is_section_input.lower() == 'д'
        
        if is_section:
            new_tasks.append({
                "type": "section",
                "num": num,
                "name": name
            })
            print("✅ Раздел добавлен.\n")
            continue
            
        unit = input("3. Единица измерения (например, м3, м2, т): ").strip()
        if unit.lower() in ['выход', 'exit', 'quit']:
            break
            
        qty_input = input("4. Количество (объем): ").strip()
        if qty_input.lower() in ['выход', 'exit', 'quit']:
            break
        try:
            qty = float(qty_input) if '.' in qty_input else int(qty_input)
        except ValueError:
            qty = qty_input
            
        crew_input = input("5. Количество человек в бригаде: ").strip()
        if crew_input.lower() in ['выход', 'exit', 'quit']:
            break
        crew = int(crew_input) if crew_input.isdigit() else 1
        
        if last_end_date:
            default_start = last_end_date + datetime.timedelta(days=1)
        else:
            default_start = datetime.date(2026, 10, 1)
            
        start_date = get_date_input_with_exit("6. Дата начала работы", default_start)
        if start_date == "exit":
            break
            
        days_exit = False
        while True:
            days_input = input("7. Длительность работы (в днях): ").strip()
            if days_input.lower() in ['выход', 'exit', 'quit']:
                days_exit = True
                break
            if days_input.isdigit() and int(days_input) > 0:
                days = int(days_input)
                break
            print("❌ Длительность должна быть целым числом больше 0!")
            
        if days_exit:
            break
            
        last_end_date = start_date + datetime.timedelta(days=days - 1)
        
        new_tasks.append({
            "type": "task",
            "num": num,
            "name": name,
            "unit": unit,
            "qty": qty,
            "crew": crew,
            "start": start_date.strftime("%Y-%m-%d"),
            "days": days
        })
        print(f"✅ Работа добавлена. Расчетное окончание: {last_end_date.strftime('%Y-%m-%d')}\n")
        
    if new_tasks:
        confirm = input("\n💾 Сохранить введенные работы и перезаписать tasks.json? (д/н): ").strip().lower()
        if confirm == 'д':
            save_tasks(new_tasks)
            print("💾 Файл tasks.json успешно обновлен!")
        else:
            print("⚠ Изменения не сохранены.")
def create_gantt_chart(output_file="Gantt_Compact.xlsx"):
    ans = input("Хотите ввести новый список строительных работ в терминале? (д/н) [Enter для н]: ").strip().lower()
    if ans == 'д':
        input_tasks_interactively()

    print("\n--- НАСТРОЙКА ДИАПАЗОНА ДИАГРАММЫ ГАНТА ---")
    # Используем обновленную функцию ввода дат с поддержкой безопасного выхода
    base_date = get_date_input_with_exit("Введите дату НАЧАЛА графика", datetime.date(2026, 10, 1))
    if base_date == "exit":
        print("🛑 Выход из программы.")
        return

    end_date = get_date_input_with_exit("Введите дату ОКОНЧАНИЯ графика", datetime.date(2027, 6, 30))
    if end_date == "exit":
        print("🛑 Выход из программы.")
        return
    
    if end_date < base_date:
        print("⚠ Дата окончания не может быть раньше даты начала! Поменял их местами.")
        base_date, end_date = end_date, base_date

    days_to_generate = (end_date - base_date).days + 1
    print(f"📊 Будет сгенерировано дней в календаре: {days_to_generate}\n")

    tasks = load_tasks()
    if not tasks:
        print("❌ Ошибка: Список задач пуст! Заполните tasks.json.")
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "График компактный"
    ws.views.sheetView[0].showGridLines = True

    HEADER_BG = "1F4E78"
    MONTH_BG = "2F5597"
    HEADER_FG = "FFFFFF"
    BORDER_COLOR = "D9D9D9"

    thin_border = Border(
        left=Side(style='thin', color=BORDER_COLOR),
        right=Side(style='thin', color=BORDER_COLOR),
        top=Side(style='thin', color=BORDER_COLOR),
        bottom=Side(style='thin', color=BORDER_COLOR)
    )

    ws["A1"] = "СОГЛАСОВАНО"
    ws["A2"] = "Должность _____________"
    ws["A3"] = "Организация ___________"
    ws["A4"] = "____/_______/ (Ф.И.О.)"
    ws["A5"] = "«___» ________ 202__ г."

    ws["J1"] = "УТВЕРЖДАЮ"
    ws["J2"] = "Должность _____________"
    ws["J3"] = "Организация ___________"
    ws["J4"] = "____/_______/ (Ф.И.О.)"
    ws["J5"] = "«___» ________ 202__ г."

    for r in range(1, 6):
        ws[f"A{r}"].font = Font(name="Calibri", size=9, bold=(r==1))
        ws[f"J{r}"].font = Font(name="Calibri", size=9, bold=(r==1))

    ws.merge_cells("A7:N7")
    title = ws["A7"]
    title.value = "ГРАФИК ПРОИЗВОДСТВА БЕТОННЫХ РАБОТ"
    title.font = Font(name="Calibri", size=14, bold=True, color="1F4E78")
    title.alignment = Alignment(horizontal="center", vertical="center")

    headers_text = ["№", "Наименование работ", "Ед. изм.", "Кол-во", "Кол-во чел.", "Начало", "Дней", "Окончание"]
    for col_idx, h_text in enumerate(headers_text, start=1):
        col_letter = get_column_letter(col_idx)
        ws.merge_cells(f"{col_letter}9:{col_letter}10")
        cell = ws.cell(row=9, column=col_idx, value=h_text)
        cell.font = Font(name="Calibri", size=10, bold=True, color=HEADER_FG)
        cell.fill = PatternFill(start_color=HEADER_BG, end_color=HEADER_BG, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.cell(row=10, column=col_idx).border = thin_border
        cell.border = thin_border

    month_names_ru = {
        1: "ЯНВАРЬ", 2: "ФЕВРАЛЬ", 3: "МАРТ", 4: "АПРЕЛЬ",
        5: "МАЙ", 6: "ИЮНЬ", 7: "ИЮЛЬ", 8: "АВГУСТ",
        9: "СЕНТЯБРЬ", 10: "ОКТЯБРЬ", 11: "НОЯБРЬ", 12: "ДЕКАБРЬ"
    }

    months_cols = {}

    for i in range(days_to_generate):
        col_idx = 9 + i
        current_date = base_date + datetime.timedelta(days=i)
        
        cell_day = ws.cell(row=10, column=col_idx, value=current_date)
        cell_day.number_format = 'DD'
        cell_day.font = Font(name="Calibri", size=9, bold=True, color=HEADER_FG)
        cell_day.fill = PatternFill(start_color=HEADER_BG, end_color=HEADER_BG, fill_type="solid")
        cell_day.alignment = Alignment(horizontal="center", vertical="center")
        cell_day.border = thin_border
        
        m_key = (current_date.year, current_date.month)
        if m_key not in months_cols:
            months_cols[m_key] = []
        months_cols[m_key].append(col_idx)

    for m_tuple, cols in months_cols.items():
        start_c = get_column_letter(cols[0])
        end_c = get_column_letter(cols[-1])
        
        ws.merge_cells(f"{start_c}9:{end_c}9")
        y_val, m_val = m_tuple
        m_text = f"{month_names_ru[m_val]} {y_val}"
        m_cell = ws.cell(row=9, column=cols[0], value=m_text)
        m_cell.font = Font(name="Calibri", size=10, bold=True, color=HEADER_FG)
        m_cell.fill = PatternFill(start_color=MONTH_BG, end_color=MONTH_BG, fill_type="solid")
        m_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        for c in cols:
            ws.cell(row=9, column=c).border = thin_border

    start_row = 11
    end_col = 8 + days_to_generate

    for row_offset, row_item in enumerate(tasks):
        current_row = start_row + row_offset
        is_section_header = (row_item.get("type") == "section")
        
        if is_section_header:
            ws.cell(row=current_row, column=1, value=row_item["num"])
            ws.cell(row=current_row, column=2, value=row_item["name"])
            for c in range(1, end_col + 1):
                cell = ws.cell(row=current_row, column=c)
                cell.font = Font(name="Calibri", size=10, bold=True)
                cell.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
                cell.border = thin_border
        else:
            ws.cell(row=current_row, column=1, value=row_item["num"])
            ws.cell(row=current_row, column=2, value=row_item["name"])
            ws.cell(row=current_row, column=3, value=row_item.get("unit", ""))
            ws.cell(row=current_row, column=4, value=row_item.get("qty", ""))
            ws.cell(row=current_row, column=5, value=row_item.get("crew", ""))
            
            d_start = datetime.datetime.strptime(row_item["start"], "%Y-%m-%d").date()
            c_start = ws.cell(row=current_row, column=6, value=d_start)
            c_start.number_format = 'DD.MM'
            
            ws.cell(row=current_row, column=7, value=row_item.get("days", 1))
            
            formula_end = f"=F{current_row}+G{current_row}-1"
            c_end = ws.cell(row=current_row, column=8, value=formula_end)
            c_end.number_format = 'DD.MM'
            
            for c in range(1, end_col + 1):
                cell = ws.cell(row=current_row, column=c)
                cell.font = Font(name="Calibri", size=10)
                cell.border = thin_border
                
                if str(c) in "1345678":
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif c == 2:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

            rule_formula = f"=AND(I$10>=$F{current_row},I$10<=$H{current_row})"
            rule = FormulaRule(
                formula=[rule_formula], 
                fill=PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
            )
            gantt_range = f"I{current_row}:{get_column_letter(end_col)}{current_row}"
            ws.conditional_formatting.add(gantt_range, rule)

    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 35
    ws.column_dimensions['C'].width = 10
    for col in ['D', 'E', 'F', 'G', 'H']:
        ws.column_dimensions[col].width = 11

    for c in range(9, end_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 3.5

    wb.save(output_file)
    print(f"🎉 Успешно! Файл {output_file} сохранен.")
